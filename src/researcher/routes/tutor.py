"""Language Tutor API routes.

Separate endpoints from the regular chat/session routes.
Manages tutor sessions, conversations, lessons, vocabulary, quizzes,
and student appraisal.
"""

import sys
import io
import re
import json
import queue
import asyncio
import threading
import logging
from pathlib import Path

from fastapi import HTTPException, Request
from fastapi.responses import StreamingResponse

from researcher.auth import get_current_user
from researcher.config import _QueueWriter, _SESSION_ID_RE
from researcher.gpu_broker import GPU, Need, queue_message
from researcher.tutor.crew import TutorCrew
from researcher.user_crews import get_crew, resolve_model
from researcher.tutor.failures import describe_for_user

from researcher.tutor.models import (
    TutorSessionCreate,
    TutorSessionUpdate,
    TutorChatRequest,
    LessonPlanRequest,
    PassageRequest,
    QuizRequest,
    QuizSubmitRequest,
    VocabularyAddRequest,
    AppraisalRequest,
)
from researcher.tutor.storage import (
    create_tutor_session,
    list_tutor_sessions,
    get_tutor_session,
    update_tutor_session_messages,
    delete_tutor_session,
    set_topic_index,
    add_vocabulary,
    add_vocabulary_batch,
    list_vocabulary,
    list_vocabulary_by_lang,
    list_vocabulary_due,
    delete_vocabulary,
    save_lesson_plan,
    list_lesson_plans,
    get_lesson_plan,
    save_quiz,
    get_quiz,
    save_quiz_results,
    list_quiz_results,
    list_quiz_results_by_lang,
    update_vocabulary_mastery,
    get_student_stats,
    get_cross_session_stats,
    save_passage,
    list_passages,
    get_passage,
    delete_passage,
    compute_unknown_ratio,
    known_words_for,
)

logger = logging.getLogger(__name__)


_CURRICULUM_FILE = Path(__file__).parent.parent / "tutor" / "curriculum.yaml"
_curriculum: dict = {}


def _curriculum_topics(level: str) -> list[str]:
    """Ordered topic list for a CEFR level, loaded once and cached."""
    global _curriculum
    if not _curriculum:
        try:
            import yaml

            with open(_CURRICULUM_FILE) as f:
                _curriculum = yaml.safe_load(f) or {}
        except Exception as exc:  # pragma: no cover - config problem
            logger.warning("Failed to load curriculum.yaml: %s", exc)
            _curriculum = {}
    return _curriculum.get((level or "A1").upper(), []) or []


def _current_curriculum_topic(session: dict) -> str:
    """The topic the learner is currently on, or "" if the level has none."""
    topics = _curriculum_topics(session.get("level", "A1"))
    if not topics:
        return ""
    idx = int(session.get("topic_index") or 0)
    return topics[min(idx, len(topics) - 1)]


async def _get_tutor_crew(request: Request, uid: str) -> TutorCrew:
    """The per-user TutorCrew, built with the model that person chose.

    Async because the choice lives in the database. It used to be read from
    app.state.current_model, which is reset from .env on restart — so the
    tutor ran whatever .env named rather than what the picker showed.
    """
    return await get_crew(request, uid, "tutor")


def _validate_sid(session_id: str):
    if not _SESSION_ID_RE.match(session_id):
        raise HTTPException(status_code=400, detail="Invalid session ID")


def _build_tutor_context(messages: list, max_messages: int = 10) -> str:
    """Build conversation context string from recent tutor messages."""
    if not messages:
        return "No previous conversation."
    recent = messages[-max_messages:]
    lines = []
    for msg in recent:
        role = msg.get("role", "user")
        text = msg.get("content", "")[:500]
        # Normalize role to lowercase for consistent formatting
        role = role.lower() if role else "user"
        lines.append(f"{role.capitalize()}: {text}")
    return "\n".join(lines)


def _sse(event: str, data) -> str:
    return f"event: {event}\ndata: {json.dumps(data)}\n\n"


def _run_tutor_crew(crew, q: queue.Queue):
    """Run a tutor crew in a thread, capturing stdout."""
    from researcher.routes.chat import _make_step_callback, _make_event_handlers
    from crewai.events.event_bus import crewai_event_bus

    crew.step_callback = _make_step_callback(q)
    handlers = _make_event_handlers(q)
    writer = _QueueWriter(q)
    old_stdout, old_stdin = sys.stdout, sys.stdin
    sys.stdout = writer
    sys.stdin = io.StringIO("N\n")
    try:
        return crew.kickoff()
    finally:
        writer.flush()
        sys.stdout = old_stdout
        sys.stdin = old_stdin
        for event_type, handler in handlers:
            crewai_event_bus.off(event_type, handler)


async def _stream_tutor_crew(crew, request: Request, uid: str, timeout_sec: int | None = None, on_complete=None):
    """Generic SSE streamer for any tutor crew run."""
    # Lazy-init per-user semaphore and cancel event
    if uid not in request.app.state.crew_semaphores:
        request.app.state.crew_semaphores[uid] = asyncio.Semaphore(1)
    if uid not in request.app.state.cancel_events:
        request.app.state.cancel_events[uid] = threading.Event()

    semaphore = request.app.state.crew_semaphores[uid]
    if timeout_sec is None:
        timeout_sec = request.app.state.crew_timeout
    q: queue.Queue = queue.Queue()

    _verbose_noise = re.compile(
        r"^(Agent:|Task:|Thought:|Action:|Action Input:|Observation:|"
        r"Entering new CrewAgentExecutor|Finished chain|"
        r"> |I encountered an error|Tool .* accepts these inputs|"
        r"Tool Name:|Tool Arguments:|Tool Description:|"
        r"\[1m|> Entering|> Finished|Moving on then)",
        re.IGNORECASE,
    )

    def _should_show(line: str) -> bool:
        if line and line[0] in "💭🔧📋✅🧠📝⚠️👁️🎯":
            return True
        if _verbose_noise.search(line):
            return False
        return True

    async def event_stream():
        cancel = request.app.state.cancel_events[uid]
        cancel.clear()
        # The per-user semaphore stops one person racing themselves. The GPU
        # broker is process-wide: one card, so the queue decides whose turn it
        # is and makes the right model resident before the work starts.
        async with semaphore:
            # Name the model that will actually run, which is the person's
            # saved choice — not the process-wide default from .env.
            _chosen, _ = await resolve_model(request, uid)
            model = Need("ollama", str(_chosen).replace("ollama/", ""))
            async with GPU.reserve(model, "tutor", user=uid) as _res:
                # Consuming wait() is what queues. Because it is an async
                # generator, the position reaches the browser while we are
                # still waiting rather than in a burst afterwards.
                async for _notice in _res.wait():
                    yield _sse("queue", queue_message(
                        _notice["position"], _notice["estimate"],
                        _notice["ticket"]))
                loop = asyncio.get_running_loop()
                future = loop.run_in_executor(None, lambda: _run_tutor_crew(crew, q))
                all_lines: list[str] = []
                elapsed = 0.0
                early_exit: str | None = None  # "cancelled" | "timeout"

                while not future.done():
                    if cancel.is_set() and early_exit is None:
                        early_exit = "cancelled"
                        yield _sse("cancelled", "Request cancelled by user")
                    if elapsed >= timeout_sec and early_exit is None:
                        early_exit = "timeout"
                        yield _sse(
                            "error",
                            f"Request timed out after {timeout_sec}s. A model that "
                            f"does not fit in VRAM runs partly on CPU and can take "
                            f"minutes per call — check `ollama ps` for a CPU/GPU "
                            f"split.")
                    await asyncio.sleep(0.15)
                    elapsed += 0.15
                    while not q.empty():
                        try:
                            line = q.get_nowait()
                            all_lines.append(line)
                            if early_exit is None and _should_show(line):
                                yield _sse("reasoning", line)
                        except queue.Empty:
                            break

                if early_exit is not None:
                    # Drain queue silently; semaphore releases after thread finished
                    while not q.empty():
                        q.get_nowait()
                    return

                # Drain remaining
                while not q.empty():
                    try:
                        line = q.get_nowait()
                        all_lines.append(line)
                        if _should_show(line):
                            yield _sse("reasoning", line)
                    except queue.Empty:
                        break

                try:
                    result = future.result()
                except Exception as e:
                    # Name the likely cause: "Invalid response from LLM call"
                    # means an exhausted token budget far more often than it means
                    # anything else, and saying so saves an evening.
                    yield _sse("error", describe_for_user(e, "generation"))
                    return

                response_text = result.raw if hasattr(result, "raw") else str(result)
                # Strip thinking tags
                response_text = re.sub(
                    r"<think>[\s\S]*?</think>\s*", "", response_text
                ).strip()

                yield _sse(
                    "done",
                    {
                        "response": response_text,
                        "reasoning": all_lines,
                    },
                )

                if on_complete is not None:
                    try:
                        await on_complete(response_text)
                    except Exception as _cb_err:
                        import logging as _log
                        _log.getLogger(__name__).warning("on_complete error: %s", _cb_err)

    return StreamingResponse(event_stream(), media_type="text/event-stream")


def register_tutor_routes(app):
    """Attach all /tutor/* endpoints to the FastAPI app."""

    # ------------------------------------------------------------------
    # Tutor Sessions CRUD
    # ------------------------------------------------------------------

    @app.get("/tutor/sessions")
    async def tutor_list_sessions(request: Request):
        user = await get_current_user(request)
        db = request.app.state.db
        sessions = await list_tutor_sessions(db, user["id"])
        return {"sessions": sessions}

    @app.post("/tutor/sessions")
    async def tutor_create_session(req: TutorSessionCreate, request: Request):
        user = await get_current_user(request)
        db = request.app.state.db
        session = await create_tutor_session(
            db, user["id"], req.name, req.target_lang, req.native_lang, req.level
        )
        return session

    @app.get("/tutor/sessions/{session_id}")
    async def tutor_get_session(session_id: str, request: Request):
        _validate_sid(session_id)
        user = await get_current_user(request)
        db = request.app.state.db
        session = await get_tutor_session(db, session_id, user["id"])
        if not session:
            raise HTTPException(status_code=404, detail="Tutor session not found")
        return session

    @app.put("/tutor/sessions/{session_id}")
    async def tutor_update_session(
        session_id: str, req: TutorSessionUpdate, request: Request
    ):
        _validate_sid(session_id)
        user = await get_current_user(request)
        db = request.app.state.db
        session = await get_tutor_session(db, session_id, user["id"])
        if not session:
            raise HTTPException(status_code=404, detail="Tutor session not found")
        await update_tutor_session_messages(db, session_id, req.messages)
        return {"id": session_id, "name": req.name}

    @app.delete("/tutor/sessions/{session_id}")
    async def tutor_delete_session(session_id: str, request: Request):
        _validate_sid(session_id)
        user = await get_current_user(request)
        db = request.app.state.db
        deleted = await delete_tutor_session(db, session_id, user["id"])
        if not deleted:
            raise HTTPException(status_code=404, detail="Tutor session not found")
        return {"deleted": session_id}

    # ------------------------------------------------------------------
    # Conversation (SSE streaming)
    # ------------------------------------------------------------------

    # ------------------------------------------------------------------
    # Scenarios (situational dialogs)
    # ------------------------------------------------------------------

    import yaml as _scenarios_yaml
    from pathlib import Path as _scenarios_path

    _SCENARIOS_PATH = _scenarios_path(__file__).parent.parent / "tutor" / "scenarios.yaml"
    _scenarios_cache: list = []

    def _load_scenarios() -> list:
        if not _scenarios_cache:
            try:
                with open(_SCENARIOS_PATH, "r", encoding="utf-8") as f:
                    _scenarios_cache.extend(_scenarios_yaml.safe_load(f) or [])
            except Exception as exc:
                logger.warning("Failed to load scenarios.yaml: %s", exc)
        return _scenarios_cache

    @app.get("/tutor/scenarios")
    async def tutor_list_scenarios(level: str = ""):
        """Return available roleplay scenarios, optionally filtered by level."""
        scenarios = _load_scenarios()
        if level:
            scenarios = [s for s in scenarios if s.get("level", "") == level]
        return {"scenarios": scenarios}

    @app.post("/tutor/chat")
    async def tutor_chat(req: TutorChatRequest, request: Request):
        """Conversational exchange with the tutor agent — SSE streaming."""
        _validate_sid(req.session_id)
        user = await get_current_user(request)
        db = request.app.state.db

        session = await get_tutor_session(db, req.session_id, user["id"])
        if not session:
            raise HTTPException(status_code=404, detail="Tutor session not found")

        uid = user["id"]
        tutor_crew = await _get_tutor_crew(request, uid)
        context = _build_tutor_context(session["messages"])

        scenario = None
        if req.mode == "scenario" and req.scenario_id:
            scenario = next(
                (s for s in _load_scenarios() if s.get("id") == req.scenario_id),
                None,
            )
            if not scenario:
                raise HTTPException(status_code=404, detail="Scenario not found")
            crew = tutor_crew.build_scenario_crew(
                message=req.message,
                context=context,
                target_lang=session["target_lang"],
                native_lang=session["native_lang"],
                level=session["level"],
                scenario=scenario,
            )
        else:
            crew = tutor_crew.build_conversation_crew(
                message=req.message,
                context=context,
                target_lang=session["target_lang"],
                native_lang=session["native_lang"],
                level=session["level"],
            )

        _scenario_ref = scenario  # capture for closure

        async def _save_chat_vocab(response_text: str):
            # Save vocabulary extracted from response
            vocab_entries = tutor_crew.extract_vocabulary_from_response(response_text)
            if vocab_entries:
                await add_vocabulary_batch(
                    db, uid, req.session_id, vocab_entries,
                    target_lang=session["target_lang"],
                )
            # Auto-save scenario transcript to lesson_plans after each turn
            if _scenario_ref and req.save_transcript:
                updated_context = _build_tutor_context(session["messages"])
                transcript = (
                    f"# Scenario: {_scenario_ref.get('title', '')}\n\n"
                    f"**Setting**: {_scenario_ref.get('setting', '')}\n"
                    f"**Your role**: {_scenario_ref.get('role_student', '')}\n\n"
                    f"---\n\n{updated_context}"
                )
                await save_lesson_plan(
                    db, uid, req.session_id,
                    topic=f"Scenario: {_scenario_ref.get('title', '')}",
                    lesson_type="scenario",
                    content={"markdown": transcript},
                )

        return await _stream_tutor_crew(crew, request, uid,
            timeout_sec=tutor_crew.get_llm_params().get("crew_timeout_sec"),
            on_complete=_save_chat_vocab)

    # ------------------------------------------------------------------
    # Lesson Plans
    # ------------------------------------------------------------------

    @app.post("/tutor/lessons")
    async def tutor_create_lesson(req: LessonPlanRequest, request: Request):
        """Generate a lesson plan — SSE streaming."""
        _validate_sid(req.session_id)
        user = await get_current_user(request)
        db = request.app.state.db

        session = await get_tutor_session(db, req.session_id, user["id"])
        if not session:
            raise HTTPException(status_code=404, detail="Tutor session not found")

        uid = user["id"]
        tutor_crew = await _get_tutor_crew(request, uid)

        crew = tutor_crew.build_lesson_crew(
            topic=req.topic,
            lesson_type=req.lesson_type,
            target_lang=session["target_lang"],
            native_lang=session["native_lang"],
            level=session["level"],
        )

        return await _stream_tutor_crew(crew, request, uid,
            timeout_sec=tutor_crew.get_llm_params().get("crew_timeout_sec"))

    @app.post("/tutor/lessons/save")
    async def tutor_save_lesson(request: Request):
        """Save a generated lesson plan to the database."""
        user = await get_current_user(request)
        db = request.app.state.db
        body = await request.json()

        session_id = body.get("session_id", "")
        _validate_sid(session_id)
        session = await get_tutor_session(db, session_id, user["id"])
        if not session:
            raise HTTPException(status_code=404, detail="Tutor session not found")

        tutor_crew = await _get_tutor_crew(request, user["id"])

        # Extract vocabulary from the lesson content and save it
        content_text = body.get("content", "")
        vocab_entries = tutor_crew.extract_vocabulary_from_response(content_text)
        if vocab_entries:
            await add_vocabulary_batch(
                db,
                user["id"],
                session_id,
                vocab_entries,
                target_lang=session["target_lang"],
            )

        lesson = await save_lesson_plan(
            db,
            user_id=user["id"],
            session_id=session_id,
            title=body.get("title", "Untitled Lesson"),
            topic=body.get("topic", ""),
            lesson_type=body.get("lesson_type", "mixed"),
            target_lang=session["target_lang"],
            level=session["level"],
            content={"markdown": content_text},
        )
        lesson["vocabulary_saved"] = len(vocab_entries)
        return lesson

    # ── Reading passages (comprehensible input) ──

    @app.post("/tutor/passages")
    async def tutor_create_passage(req: PassageRequest, request: Request):
        """Generate a levelled reading passage — SSE streaming.

        Seeds the passage with vocabulary that is due for review, so the learner
        meets those words in context instead of as isolated flashcards.
        """
        _validate_sid(req.session_id)
        user = await get_current_user(request)
        db = request.app.state.db

        session = await get_tutor_session(db, req.session_id, user["id"])
        if not session:
            raise HTTPException(status_code=404, detail="Tutor session not found")

        due: list[dict] = []
        if req.seed_due_vocab:
            due = await list_vocabulary_due(
                db, user["id"], session["target_lang"]
            )
            due = due[: max(0, req.max_seed_words)]

        topic = req.topic.strip()
        if not topic:
            topic = _current_curriculum_topic(session) or "everyday life"

        uid = user["id"]
        tutor_crew = await _get_tutor_crew(request, uid)
        crew = tutor_crew.build_passage_crew(
            topic=topic,
            target_lang=session["target_lang"],
            native_lang=session["native_lang"],
            level=session["level"],
            due_words=due,
            word_count=req.word_count,
        )
        return await _stream_tutor_crew(
            crew, request, uid,
            timeout_sec=tutor_crew.get_llm_params().get("crew_timeout_sec"),
        )

    @app.post("/tutor/passages/save")
    async def tutor_save_passage(request: Request):
        """Persist a generated passage, measuring how much of it is new.

        The unknown-word ratio is what tells the learner whether the passage is
        pitched at i+1 (~10% new). It is computed here rather than trusted from
        the model, which has no idea what this particular learner knows.
        """
        user = await get_current_user(request)
        body = await request.json()
        session_id = (body.get("session_id") or "").strip()
        _validate_sid(session_id)
        content = (body.get("content") or "").strip()
        if not content:
            raise HTTPException(status_code=400, detail="content is required")

        db = request.app.state.db
        session = await get_tutor_session(db, session_id, user["id"])
        if not session:
            raise HTTPException(status_code=404, detail="Tutor session not found")

        known = await known_words_for(db, user["id"], session["target_lang"])
        ratio, word_count, unknown_sample = compute_unknown_ratio(
            content, known, session["target_lang"]
        )

        title = (body.get("title") or "").strip()
        if not title:
            for line in content.splitlines():
                if line.startswith("#"):
                    title = line.lstrip("#").strip()
                    break
            title = title or "Reading passage"

        saved = await save_passage(
            db, user["id"], session_id, title,
            body.get("topic", ""), session["target_lang"], session["level"],
            content, body.get("seeded_words", []), word_count, ratio,
        )
        saved["unknown_sample"] = unknown_sample
        # The ratio is only meaningful once the learner's saved vocabulary is
        # large enough to represent what they actually know. Send the size so
        # the UI can say "can't judge yet" instead of always reporting "hard".
        saved["known_word_count"] = len(known)
        return saved

    @app.get("/tutor/sessions/{session_id}/passages")
    async def tutor_list_passages(session_id: str, request: Request):
        _validate_sid(session_id)
        user = await get_current_user(request)
        return {
            "passages": await list_passages(
                request.app.state.db, user["id"], session_id
            )
        }

    @app.get("/tutor/passages/{passage_id}")
    async def tutor_get_passage(passage_id: str, request: Request):
        user = await get_current_user(request)
        p = await get_passage(request.app.state.db, user["id"], passage_id)
        if not p:
            raise HTTPException(status_code=404, detail="Passage not found")
        return p

    @app.delete("/tutor/passages/{passage_id}")
    async def tutor_delete_passage(passage_id: str, request: Request):
        user = await get_current_user(request)
        ok = await delete_passage(request.app.state.db, user["id"], passage_id)
        if not ok:
            raise HTTPException(status_code=404, detail="Passage not found")
        return {"deleted": passage_id}

    @app.get("/tutor/sessions/{session_id}/lessons")
    async def tutor_list_lessons(session_id: str, request: Request):
        _validate_sid(session_id)
        user = await get_current_user(request)
        db = request.app.state.db
        lessons = await list_lesson_plans(db, user["id"], session_id)
        return {"lessons": lessons}

    @app.get("/tutor/lessons/{lesson_id}")
    async def tutor_get_lesson(lesson_id: str, request: Request):
        user = await get_current_user(request)
        db = request.app.state.db
        lesson = await get_lesson_plan(db, lesson_id, user["id"])
        if not lesson:
            raise HTTPException(status_code=404, detail="Lesson not found")
        return lesson

    @app.get("/tutor/lessons/{lesson_id}/export")
    async def tutor_export_lesson(
        lesson_id: str, request: Request, format: str = "md"
    ):
        """Export a lesson plan as Markdown or DOCX."""
        from fastapi.responses import Response

        user = await get_current_user(request)
        db = request.app.state.db
        lesson = await get_lesson_plan(db, lesson_id, user["id"])
        if not lesson:
            raise HTTPException(status_code=404, detail="Lesson not found")

        content = lesson.get("content", {})
        title = lesson.get("title", "Lesson")
        lang = lesson.get("target_lang", "")
        level = lesson.get("level", "")

        # Content is stored as {"markdown": text} — unwrap it
        md_text = content.get("markdown", "") if isinstance(content, dict) else str(content)

        safe = "".join(c if c.isalnum() or c in "-_ " else "_" for c in title).strip()

        if format == "md":
            return Response(
                content=md_text.encode("utf-8"),
                media_type="text/markdown",
                headers={"Content-Disposition": f'attachment; filename="{safe}.md"'},
            )

        if format == "docx":
            try:
                from docx import Document
                from docx.shared import Pt, RGBColor
                import io as _io

                doc = Document()
                doc.add_heading(title, 0)
                doc.add_paragraph(f"Language: {lang}  |  Level: {level}  |  Topic: {lesson.get('topic', '')}")

                # Parse markdown line-by-line into DOCX
                for line in md_text.splitlines():
                    if line.startswith("### "):
                        doc.add_heading(line[4:], level=3)
                    elif line.startswith("## "):
                        doc.add_heading(line[3:], level=2)
                    elif line.startswith("# "):
                        doc.add_heading(line[2:], level=1)
                    elif line.startswith("- ") or line.startswith("* "):
                        doc.add_paragraph(line[2:], style="List Bullet")
                    elif line.startswith("|"):
                        # Markdown table row — add as plain paragraph
                        cells = [c.strip() for c in line.strip("|").split("|")]
                        doc.add_paragraph("  |  ".join(cells))
                    elif line.strip():
                        doc.add_paragraph(line)

                buf = _io.BytesIO()
                doc.save(buf)
                buf.seek(0)
                return Response(
                    content=buf.read(),
                    media_type="application/vnd.openxmlformats-officedocument"
                    ".wordprocessingml.document",
                    headers={"Content-Disposition": f'attachment; filename="{safe}.docx"'},
                )
            except ImportError:
                raise HTTPException(
                    status_code=400, detail="DOCX export requires python-docx"
                )

        raise HTTPException(status_code=400, detail="format must be 'md' or 'docx'")

    # ------------------------------------------------------------------
    # Vocabulary
    # ------------------------------------------------------------------

    @app.get("/tutor/sessions/{session_id}/vocabulary")
    async def tutor_list_vocabulary(session_id: str, request: Request):
        _validate_sid(session_id)
        user = await get_current_user(request)
        db = request.app.state.db
        vocab = await list_vocabulary(db, user["id"], session_id)
        return {"vocabulary": vocab}

    @app.get("/tutor/vocabulary")
    async def tutor_list_vocabulary_by_lang(request: Request, lang: str = ""):
        """List vocabulary for a target language (cross-session)."""
        user = await get_current_user(request)
        db = request.app.state.db
        if not lang:
            raise HTTPException(status_code=400, detail="lang query parameter required")
        vocab = await list_vocabulary_by_lang(db, user["id"], lang)
        return {"vocabulary": vocab}

    @app.get("/tutor/vocabulary/due")
    async def tutor_vocabulary_due(request: Request, lang: str = ""):
        """Return vocabulary items due for SRS review today (SM-2 schedule)."""
        user = await get_current_user(request)
        db = request.app.state.db
        due = await list_vocabulary_due(db, user["id"], target_lang=lang)
        return {"due": due, "count": len(due)}

    @app.post("/tutor/vocabulary")
    async def tutor_add_vocabulary(req: VocabularyAddRequest, request: Request):
        _validate_sid(req.session_id)
        user = await get_current_user(request)
        db = request.app.state.db
        session = await get_tutor_session(db, req.session_id, user["id"])
        if not session:
            raise HTTPException(status_code=404, detail="Tutor session not found")
        target_lang = req.target_lang or session["target_lang"]
        entry = await add_vocabulary(
            db,
            user["id"],
            req.session_id,
            req.word,
            req.translation,
            req.context,
            req.phonetic,
            req.part_of_speech,
            target_lang=target_lang,
        )
        return entry

    @app.delete("/tutor/vocabulary/{vocab_id}")
    async def tutor_delete_vocabulary(vocab_id: str, request: Request):
        user = await get_current_user(request)
        db = request.app.state.db
        deleted = await delete_vocabulary(db, vocab_id, user["id"])
        if not deleted:
            raise HTTPException(status_code=404, detail="Vocabulary entry not found")
        return {"deleted": vocab_id}

    # ------------------------------------------------------------------
    # Quizzes
    # ------------------------------------------------------------------

    @app.post("/tutor/quiz/generate")
    async def tutor_generate_quiz(req: QuizRequest, request: Request):
        """Generate a quiz — SSE streaming."""
        _validate_sid(req.session_id)
        user = await get_current_user(request)
        db = request.app.state.db

        session = await get_tutor_session(db, req.session_id, user["id"])
        if not session:
            raise HTTPException(status_code=404, detail="Tutor session not found")

        uid = user["id"]
        tutor_crew = await _get_tutor_crew(request, uid)

        # Build vocabulary context from the student's saved words
        vocab = await list_vocabulary(db, user["id"], req.session_id)
        vocab_ctx = ""
        if vocab:
            vocab_lines = [f"- {v['word']} = {v['translation']}" for v in vocab[:50]]
            vocab_ctx = (
                "The student has learned these vocabulary words.\n"
                "Include some of them in the quiz:\n" + "\n".join(vocab_lines)
            )

        # Build lesson context if scoped to a lesson
        lesson_ctx = ""
        if req.lesson_id:
            lesson = await get_lesson_plan(db, req.lesson_id, user["id"])
            if lesson and lesson.get("content"):
                content = lesson["content"]
                md = (
                    content.get("markdown", "")[:2000]
                    if isinstance(content, dict)
                    else str(content)[:2000]
                )
                lesson_ctx = f"Base the quiz on this lesson:\n{md}"

        crew = tutor_crew.build_quiz_crew(
            quiz_type=req.quiz_type,
            num_questions=req.num_questions,
            target_lang=session["target_lang"],
            native_lang=session["native_lang"],
            level=session["level"],
            vocabulary_context=vocab_ctx,
            lesson_context=lesson_ctx,
        )

        return await _stream_tutor_crew(crew, request, uid,
            timeout_sec=tutor_crew.get_llm_params().get("crew_timeout_sec"))

    @app.post("/tutor/quiz/save")
    async def tutor_save_quiz(request: Request):
        """Save a generated quiz (the JSON questions) for later answering."""
        user = await get_current_user(request)
        db = request.app.state.db
        body = await request.json()

        session_id = body.get("session_id", "")
        _validate_sid(session_id)

        tutor_crew = await _get_tutor_crew(request, user["id"])
        response_text = body.get("response", "")
        questions = tutor_crew.extract_quiz_json(response_text)
        if not questions:
            raise HTTPException(
                status_code=400, detail="Could not parse quiz JSON from response"
            )

        quiz_type = body.get("quiz_type", "mixed")
        questions = tutor_crew.filter_quiz_questions(questions, quiz_type)

        quiz = await save_quiz(
            db,
            user_id=user["id"],
            session_id=session_id,
            quiz_type=quiz_type,
            questions=questions,
            lesson_id=body.get("lesson_id"),
        )
        quiz["questions"] = questions
        return quiz

    @app.post("/tutor/quiz/submit")
    async def tutor_submit_quiz(req: QuizSubmitRequest, request: Request):
        """Submit answers for a quiz, grade them, and return feedback."""
        _validate_sid(req.session_id)
        user = await get_current_user(request)
        db = request.app.state.db

        quiz = await get_quiz(db, req.quiz_id, user["id"])
        if not quiz:
            raise HTTPException(status_code=404, detail="Quiz not found")

        questions = quiz["questions"]
        total = len(questions)

        # Fetch session language info for semantic grading
        session = await get_tutor_session(db, req.session_id, user["id"])
        target_lang = session["target_lang"] if session else "English"
        native_lang = session["native_lang"] if session else "English"
        level = session["level"] if session else "A1"

        # Split: pronunciation answers have pre-computed scores; the rest go to LLM
        pre_graded: dict[int, dict] = {}
        llm_questions = []
        llm_answers = []
        for q, ans in zip(questions, req.answers):
            qid = q.get("id", questions.index(q))
            # Blank answer — never correct, skip LLM entirely
            if not (ans.get("answer") or "").strip():
                pre_graded[qid] = {
                    "question_id": qid,
                    "score": 0.0,
                    "is_correct": False,
                    "feedback": "No answer provided.",
                    "correct_answer": q.get("correct_answer", ""),
                }
                continue
            if q.get("type") == "pronunciation":
                try:
                    score_val = float(ans.get("answer", "0") or "0")
                except (ValueError, TypeError):
                    score_val = 0.0
                score_val = max(0.0, min(1.0, score_val))
                pre_graded[qid] = {
                    "question_id": qid,
                    "score": score_val,
                    "is_correct": score_val >= 0.5,
                    "feedback": "Pronunciation scored automatically.",
                    "correct_answer": q.get("correct_answer", ""),
                }
            else:
                llm_questions.append(q)
                llm_answers.append(ans)

        tutor_crew = await _get_tutor_crew(request, user["id"])
        if llm_questions:
            llm_graded = await tutor_crew.grade_answers(
                questions=llm_questions,
                student_answers=llm_answers,
                target_lang=target_lang,
                native_lang=native_lang,
                level=level,
            )
        else:
            llm_graded = []

        # Merge: restore original question order
        graded_map = {g["question_id"]: g for g in llm_graded}
        graded_map.update(pre_graded)
        graded = [graded_map[q.get("id", i)] for i, q in enumerate(questions) if q.get("id", i) in graded_map]
        if len(graded) < len(questions):
            # Fallback: fill any gaps
            present = {g["question_id"] for g in graded}
            for i, q in enumerate(questions):
                qid = q.get("id", i)
                if qid not in present:
                    graded.append({"question_id": qid, "score": 0.0, "is_correct": False, "feedback": ""})

        # Compute integer score (questions with score >= 0.5) and float sum
        score = sum(1 for g in graded if g.get("is_correct"))
        score_float = sum(g.get("score", 0.0) for g in graded)

        pct = round(score_float / total * 100, 1) if total > 0 else 0
        feedback = f"Score: {score_float:.1f}/{total} ({pct}%)"

        await save_quiz_results(
            db,
            req.quiz_id,
            user["id"],
            graded,
            score,
            feedback,
            score_float=score_float,
        )

        # Update vocabulary mastery based on quiz answers
        vocab = await list_vocabulary(db, user["id"], req.session_id)
        vocab_map = {v["word"].lower(): v["id"] for v in vocab}
        for g in graded:
            qid = g["question_id"]
            if 0 <= qid < len(questions):
                q = questions[qid]
                # Try to find the word in vocabulary to update mastery
                q_text = q.get("question", "").lower()
                for word, vid in vocab_map.items():
                    if word in q_text or word in g.get("correct_answer", "").lower():
                        await update_vocabulary_mastery(db, vid, g["is_correct"])
                        break

        return {
            "quiz_id": req.quiz_id,
            "score": score_float,
            "total": total,
            "percentage": pct,
            "feedback": feedback,
            "details": graded,
        }

    @app.get("/tutor/sessions/{session_id}/quizzes")
    async def tutor_list_quizzes(session_id: str, request: Request):
        _validate_sid(session_id)
        user = await get_current_user(request)
        db = request.app.state.db
        quizzes = await list_quiz_results(db, user["id"], session_id)
        return {"quizzes": quizzes}

    @app.get("/tutor/quizzes")
    async def tutor_list_quizzes_by_lang(request: Request, lang: str = ""):
        """List all quiz results for a target language (cross-session)."""
        user = await get_current_user(request)
        db = request.app.state.db
        if not lang:
            raise HTTPException(status_code=400, detail="lang query parameter required")
        quizzes = await list_quiz_results_by_lang(db, user["id"], lang)
        return {"quizzes": quizzes}

    @app.get("/tutor/quiz/{quiz_id}")
    async def tutor_get_quiz(quiz_id: str, request: Request):
        user = await get_current_user(request)
        db = request.app.state.db
        quiz = await get_quiz(db, quiz_id, user["id"])
        if not quiz:
            raise HTTPException(status_code=404, detail="Quiz not found")
        return quiz

    # ------------------------------------------------------------------
    # Appraisal
    # ------------------------------------------------------------------

    @app.post("/tutor/appraisal")
    async def tutor_appraisal(req: AppraisalRequest, request: Request):
        """Generate a student progress appraisal — SSE streaming."""
        _validate_sid(req.session_id)
        user = await get_current_user(request)
        db = request.app.state.db

        session = await get_tutor_session(db, req.session_id, user["id"])
        if not session:
            raise HTTPException(status_code=404, detail="Tutor session not found")

        # Current session stats
        stats = await get_student_stats(db, user["id"], req.session_id)

        # Cross-session stats for same target language
        cross = await get_cross_session_stats(db, user["id"], session["target_lang"])
        stats["cross_session"] = cross

        stats_str = json.dumps(stats, indent=2)

        # Get recent messages for context
        recent = _build_tutor_context(session["messages"], max_messages=20)

        uid = user["id"]
        tutor_crew = await _get_tutor_crew(request, uid)

        crew = tutor_crew.build_appraisal_crew(
            target_lang=session["target_lang"],
            native_lang=session["native_lang"],
            level=session["level"],
            stats=stats_str,
            recent_messages=recent,
        )

        return await _stream_tutor_crew(crew, request, uid,
            timeout_sec=tutor_crew.get_llm_params().get("crew_timeout_sec"))

    # ------------------------------------------------------------------
    # Student Stats
    # ------------------------------------------------------------------

    @app.get("/tutor/sessions/{session_id}/stats")
    async def tutor_stats(session_id: str, request: Request):
        _validate_sid(session_id)
        user = await get_current_user(request)
        db = request.app.state.db
        session = await get_tutor_session(db, session_id, user["id"])
        if not session:
            raise HTTPException(status_code=404, detail="Tutor session not found")
        stats = await get_student_stats(db, user["id"], session_id)
        stats["session"] = {
            "target_lang": session["target_lang"],
            "native_lang": session["native_lang"],
            "level": session["level"],
        }

        # Mastery distribution for histogram (levels 0-5)
        cursor = await db.execute(
            "SELECT mastery_level, COUNT(*) FROM vocabulary "
            "WHERE user_id = ? AND session_id = ? GROUP BY mastery_level ORDER BY mastery_level",
            (user["id"], session_id),
        )
        mastery_rows = await cursor.fetchall()
        mastery_dist = {str(i): 0 for i in range(6)}
        for row in mastery_rows:
            mastery_dist[str(row[0])] = row[1]
        stats["mastery_distribution"] = mastery_dist

        # Recent quiz scores (last 10) for trend chart
        cursor = await db.execute(
            "SELECT score_float, total, created_at FROM quiz_results "
            "WHERE user_id = ? AND session_id = ? AND total > 0 "
            "ORDER BY created_at DESC LIMIT 10",
            (user["id"], session_id),
        )
        qrows = await cursor.fetchall()
        stats["quiz_history"] = [
            {
                "pct": round(r[0] / r[1] * 100, 1) if r[1] else 0,
                "date": r[2][:10],
            }
            for r in reversed(qrows)
        ]

        # SRS due count
        from researcher.tutor.storage import list_vocabulary_due
        due = await list_vocabulary_due(db, user["id"], session["target_lang"])
        stats["srs_due_count"] = len(due)

        return stats

    # ------------------------------------------------------------------
    # Translator
    # ------------------------------------------------------------------

    @app.post("/tutor/translate")
    async def tutor_translate(request: Request):
        """Translate text between target and native language using the LLM."""
        user = await get_current_user(request)
        body = await request.json()
        text = (body.get("text") or "").strip()
        source_lang = body.get("source_lang", "").strip()
        target_lang = body.get("target_lang", "").strip()

        if not text:
            raise HTTPException(status_code=400, detail="Empty text")
        if not source_lang or not target_lang:
            raise HTTPException(
                status_code=400, detail="source_lang and target_lang required"
            )

        tutor_crew = await _get_tutor_crew(request, user["id"])

        prompt = (
            f"Translate the following text from {source_lang} to {target_lang}.\n"
            f"Return ONLY the translation, nothing else — no explanations, "
            f"no quotation marks, no prefixes.\n\n"
            f"{text}"
        )

        loop = asyncio.get_running_loop()

        def _call_llm():
            from litellm import completion

            resp = completion(
                model=tutor_crew._model_name,
                messages=[{"role": "user", "content": prompt}],
                temperature=0.2,
                api_base="http://localhost:11434",
            )
            return resp.choices[0].message.content.strip()

        try:
            _chosen, _ = await resolve_model(request, user["id"])
            async with GPU.exclusive(
                    Need("ollama", str(_chosen).replace("ollama/", "")),
                    "translate", user=user["id"]):
                translation = await loop.run_in_executor(None, _call_llm)
            # Strip thinking tags if present (qwen3/deepseek-r1)
            translation = re.sub(
                r"<think>.*?</think>", "", translation, flags=re.DOTALL
            ).strip()
            return {
                "translation": translation,
                "source_lang": source_lang,
                "target_lang": target_lang,
            }
        except Exception as e:
            logger.error("Translation failed: %s", e)
            raise HTTPException(status_code=500, detail=f"Translation failed: {e}")

    # ------------------------------------------------------------------
    # Curriculum (CEFR topic sequence)
    # ------------------------------------------------------------------

    import yaml as _yaml
    from pathlib import Path as _Path

    _CURRICULUM_PATH = _Path(__file__).parent.parent / "tutor" / "curriculum.yaml"
    _curriculum_cache: dict = {}

    def _load_curriculum() -> dict:
        if not _curriculum_cache:
            try:
                with open(_CURRICULUM_PATH, "r", encoding="utf-8") as f:
                    _curriculum_cache.update(_yaml.safe_load(f))
            except Exception as exc:
                logger.warning("Failed to load curriculum.yaml: %s", exc)
        return _curriculum_cache

    @app.get("/tutor/curriculum")
    async def tutor_curriculum():
        """Return the full CEFR curriculum structure."""
        return {"curriculum": _load_curriculum()}

    @app.get("/tutor/sessions/{session_id}/curriculum-progress")
    async def tutor_curriculum_progress(session_id: str, request: Request):
        """Return current topic and progress within the session's level."""
        _validate_sid(session_id)
        user = await get_current_user(request)
        db = request.app.state.db
        session = await get_tutor_session(db, session_id, user["id"])
        if not session:
            raise HTTPException(status_code=404, detail="Tutor session not found")
        level = session.get("level", "A1")
        topic_index = session.get("topic_index", 0)
        curriculum = _load_curriculum()
        topics = curriculum.get(level, [])
        current_topic = topics[topic_index] if topic_index < len(topics) else None
        return {
            "level": level,
            "topic_index": topic_index,
            "total_topics": len(topics),
            "current_topic": current_topic,
            "topics": topics,
            "pct": round(topic_index / len(topics) * 100) if topics else 0,
        }

    @app.post("/tutor/sessions/{session_id}/advance-topic")
    async def tutor_advance_topic(session_id: str, request: Request):
        """Mark the current topic complete and advance to the next one."""
        _validate_sid(session_id)
        user = await get_current_user(request)
        db = request.app.state.db
        session = await get_tutor_session(db, session_id, user["id"])
        if not session:
            raise HTTPException(status_code=404, detail="Tutor session not found")
        level = session.get("level", "A1")
        curriculum = _load_curriculum()
        topics = curriculum.get(level, [])
        new_index = min(session.get("topic_index", 0) + 1, max(0, len(topics) - 1))
        await set_topic_index(db, session_id, new_index)
        current_topic = topics[new_index] if new_index < len(topics) else None
        return {
            "topic_index": new_index,
            "total_topics": len(topics),
            "current_topic": current_topic,
            "pct": round(new_index / len(topics) * 100) if topics else 0,
        }

    @app.post("/tutor/sessions/{session_id}/prev-topic")
    async def tutor_prev_topic(session_id: str, request: Request):
        """Go back to the previous curriculum topic."""
        _validate_sid(session_id)
        user = await get_current_user(request)
        db = request.app.state.db
        session = await get_tutor_session(db, session_id, user["id"])
        if not session:
            raise HTTPException(status_code=404, detail="Tutor session not found")
        level = session.get("level", "A1")
        curriculum = _load_curriculum()
        topics = curriculum.get(level, [])
        new_index = max(0, session.get("topic_index", 0) - 1)
        await set_topic_index(db, session_id, new_index)
        current_topic = topics[new_index] if new_index < len(topics) else None
        return {
            "topic_index": new_index,
            "total_topics": len(topics),
            "current_topic": current_topic,
            "pct": round(new_index / len(topics) * 100) if topics else 0,
        }
